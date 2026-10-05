"""Disposable workflow/connector fixtures; remote requests are substituted."""

import asyncio
import json
from datetime import datetime, timezone
import pytest
from jarvis.store import Store
from jarvis.workflows import Workflows, WorkflowDefinition, Schedule, next_due
from jarvis.config import Settings
from jarvis.tools import Tools
from jarvis.jobs import Jobs


class Model:
    async def stream(self, messages, tools, settings):
        assert settings.provider == "ollama"
        assert not tools
        yield {
            "message": {"content": "Verified fixture report. The garden gate is blue."}
        }


@pytest.mark.asyncio
async def test_research_monitor_stays_quiet_until_sources_change(tmp_path, monkeypatch):
    from jarvis import config

    monkeypatch.setattr(config, "DATA", tmp_path)
    store = Store(tmp_path / "db")
    jobs = Jobs(store)
    settings = Settings()
    tools = Tools(store, lambda: settings)
    observed = {"items": [{"title": "First source", "url": "https://example.com"}]}

    async def fixture_tool(*args, **kwargs):
        return {"ok": True, "data": observed}

    monkeypatch.setattr(tools, "execute", fixture_tool)
    engine = Workflows(store, jobs, tools, Model(), lambda: settings)
    definition = engine.save(
        WorkflowDefinition(title="Research", template="research_monitor")
    )
    first = engine.start(definition["id"], "research-first")
    await engine.execute(first["run_id"])
    second = engine.start(definition["id"], "research-second")
    await engine.execute(second["run_id"])
    assert json.loads(engine.run(second["run_id"])["result"])["status"] == "unchanged"
    assert len(store.all("SELECT * FROM generated_documents")) == 1
    assert len(store.all("SELECT * FROM workflow_notifications")) == 1
    observed["items"][0]["title"] = "Changed source"
    third = engine.start(definition["id"], "research-third")
    await engine.execute(third["run_id"])
    assert engine.run(third["run_id"])["state"] == "succeeded"
    assert len(store.all("SELECT * FROM generated_documents")) == 2
    assert len(store.all("SELECT * FROM workflow_notifications")) == 2
    await jobs.close()
    store.close()


@pytest.mark.asyncio
async def test_workflow_checkpoint_idempotency_and_exports(tmp_path, monkeypatch):
    from jarvis import config

    monkeypatch.setattr(config, "DATA", tmp_path)
    s = Store(tmp_path / "db.sqlite")
    jobs = Jobs(s)
    settings = Settings(provider="compatible")
    engine = Workflows(s, jobs, Tools(s, lambda: settings), Model(), lambda: settings)
    definition = WorkflowDefinition(
        title="Fixture",
        template="voice_document",
        inputs={"content": "A true fixture fact.", "format": "docx"},
    )
    record = engine.save(definition)
    run = engine.start(record["id"], "fixture-run-key")
    again = engine.start(record["id"], "fixture-run-key")
    assert run["run_id"] == again["run_id"]
    await engine.execute(run["run_id"])
    finished = engine.run(run["run_id"])
    assert finished["state"] == "succeeded"
    assert len(s.all("SELECT * FROM generated_documents")) == 1
    await engine.execute(run["run_id"])
    assert len(s.all("SELECT * FROM generated_documents")) == 1
    assert s.all("SELECT * FROM workflow_notifications")
    s.close()
    await jobs.close()


def test_overdue_schedules_coalesce_and_timezone():
    instant = datetime(2026, 10, 5, 4, 0, tzinfo=timezone.utc)
    schedule = Schedule(kind="daily", time="10:00", timezone="Asia/Kolkata")
    assert next_due(schedule, instant) == "2026-10-05T04:30:00+00:00"
    assert (
        next_due(schedule, datetime(2026, 10, 8, 6, 0, tzinfo=timezone.utc))
        == "2026-10-09T04:30:00+00:00"
    )


@pytest.mark.asyncio
async def test_startup_briefing_deduplicates_restart_and_schedule(tmp_path):
    store = Store(tmp_path / "db")
    jobs = Jobs(store)
    settings = Settings()
    engine = Workflows(
        store, jobs, Tools(store, lambda: settings), Model(), lambda: settings
    )
    definition = engine.save(
        WorkflowDefinition(
            title="Daily update",
            template="morning_briefing",
            enabled=True,
            run_on_startup=True,
            schedule={"kind": "daily", "time": "08:00"},
        )
    )
    instant = datetime(2026, 10, 5, 1, 0, tzinfo=timezone.utc)
    engine.startup(instant)
    engine.startup(instant)
    first = store.all("SELECT * FROM workflow_runs")
    assert len(first) == 1 and first[0]["trigger_slot"] == "briefing:2026-10-05"
    store.run(
        "UPDATE workflows SET next_due=? WHERE id=?",
        ("2026-10-05T02:30:00+00:00", definition["id"]),
    )
    engine.tick(datetime(2026, 10, 5, 4, 0, tzinfo=timezone.utc))
    assert len(store.all("SELECT * FROM workflow_runs")) == 1
    # A restart marks uncertain work interrupted and must not silently retry it.
    restarted = Workflows(
        store, jobs, Tools(store, lambda: settings), Model(), lambda: settings
    )
    restarted.startup(instant)
    assert restarted.run(first[0]["id"])["state"] == "interrupted"
    assert len(store.all("SELECT * FROM workflow_runs")) == 1
    restarted.startup(datetime(2026, 10, 5, 19, 0, tzinfo=timezone.utc))
    assert (
        len(store.all("SELECT * FROM workflow_runs")) == 2
    )  # already Oct 6 in Kolkata
    engine.save(
        WorkflowDefinition(
            title="Paused",
            template="morning_briefing",
            enabled=False,
            run_on_startup=True,
        )
    )
    restarted.startup(datetime(2026, 10, 5, 19, 0, tzinfo=timezone.utc))
    assert len(store.all("SELECT * FROM workflow_runs")) == 2
    await jobs.close()
    store.close()


def test_startup_cannot_launch_action_workflows():
    with pytest.raises(ValueError, match="read-only"):
        WorkflowDefinition(
            title="Unsafe startup", template="desktop_routine", run_on_startup=True
        )


@pytest.mark.asyncio
async def test_restart_does_not_replay_uncertain_work(tmp_path):
    s = Store(tmp_path / "db")
    jobs = Jobs(s)
    settings = Settings()
    w = Workflows(s, jobs, Tools(s, lambda: settings), Model(), lambda: settings)
    r = w.save(
        WorkflowDefinition(
            title="Fixture", template="voice_document", inputs={"content": "Text"}
        )
    )
    run = w.start(r["id"], "restart-fixture-key")
    s.run("UPDATE workflow_runs SET state='running' WHERE id=?", (run["run_id"],))
    s.close()
    await jobs.close()
    s = Store(tmp_path / "db")
    jobs = Jobs(s)
    w = Workflows(s, jobs, Tools(s, lambda: settings), Model(), lambda: settings)
    assert w.run(run["run_id"])["state"] == "interrupted"
    assert not s.all("SELECT * FROM generated_documents")
    s.close()
    await jobs.close()


@pytest.mark.asyncio
async def test_voice_preempts_only_readonly_generation(tmp_path):
    started = asyncio.Event()

    class Slow(Model):
        calls = 0

        async def stream(self, messages, tools, settings):
            self.calls += 1
            started.set()
            if self.calls == 1:
                await asyncio.sleep(30)
            yield {"message": {"content": "Safe resumed analysis."}}

    s = Store(tmp_path / "db")
    jobs = Jobs(s)
    model = Slow()
    settings = Settings()
    w = Workflows(s, jobs, Tools(s, lambda: settings), model, lambda: settings)
    task = asyncio.create_task(w.generate("Fixture", {}))
    await started.wait()
    w.voice_started()
    await asyncio.sleep(0.01)
    assert not task.done()
    w.voice_finished()
    assert await task == "Safe resumed analysis."
    assert model.calls == 2
    s.close()
    await jobs.close()


@pytest.mark.asyncio
async def test_cancelled_run_cannot_emit_notification(tmp_path):
    s = Store(tmp_path / "db")
    jobs = Jobs(s)
    settings = Settings()
    w = Workflows(s, jobs, Tools(s, lambda: settings), Model(), lambda: settings)
    r = w.save(
        WorkflowDefinition(
            title="Fixture", template="voice_document", inputs={"content": "Text"}
        )
    )
    run = w.start(r["id"], "cancel-fixture-key")
    w.cancel(run["run_id"])
    await w.execute(run["run_id"])
    assert w.run(run["run_id"])["state"] == "cancelled"
    assert not s.all("SELECT * FROM workflow_notifications")
    s.close()
    await jobs.close()


@pytest.mark.asyncio
async def test_router_honors_private_per_request_provider(monkeypatch):
    from jarvis.providers import ModelRouter

    router = ModelRouter(lambda: Settings(provider="compatible"))

    async def local(*args):
        yield {"message": {"content": "local only"}}

    def cloud(*args):
        pytest.fail("Private workflow routed to hosted provider")

    monkeypatch.setattr(router.ollama, "stream", local)
    monkeypatch.setattr(router.compatible, "stream", cloud)
    chunks = [x async for x in router.stream([], [], Settings(provider="ollama"))]
    assert chunks[0]["message"]["content"] == "local only"
    await router.close()


@pytest.mark.parametrize("format", ["csv", "xlsx", "pptx"])
def test_real_additional_exports(tmp_path, monkeypatch, format):
    from jarvis import config
    from jarvis.exports import create_document, document_path

    monkeypatch.setattr(config, "DATA", tmp_path)
    s = Store(tmp_path / "db")
    result = create_document(
        s, "Fixture", 'Name,Value\nRoadmap,42\nUnsafe,=HYPERLINK("bad")', format
    )
    file = document_path(s, result["id"])
    if format == "xlsx":
        from openpyxl import load_workbook

        sheet = load_workbook(file).active
        assert sheet["B2"].value == "42"
        assert sheet["B3"].data_type != "f"
    elif format == "pptx":
        from pptx import Presentation

        assert len(Presentation(file).slides) == 2
    else:
        assert "'=HYPERLINK" in file.read_text()
    assert result["status"] == "completed"
    s.close()


@pytest.mark.asyncio
async def test_worker_remains_available_after_workflow_cancel(tmp_path):
    started = asyncio.Event()

    class Slow(Model):
        calls = 0

        async def stream(self, *args):
            self.calls += 1
            if self.calls == 1:
                started.set()
                await asyncio.sleep(30)
            yield {"message": {"content": "Second result."}}

    s = Store(tmp_path / "db")
    j = Jobs(s)
    settings = Settings()
    w = Workflows(s, j, Tools(s, lambda: settings), Slow(), lambda: settings)
    d = w.save(
        WorkflowDefinition(
            title="Fixture",
            template="voice_document",
            inputs={"content": "Text", "format": "md"},
        )
    )
    j.start()
    r = w.start(d["id"])
    await started.wait()
    w.cancel(r["run_id"])
    r2 = w.start(d["id"])
    await asyncio.wait_for(j.queue.join(), 3)
    assert w.run(r["run_id"])["state"] == "cancelled"
    assert w.run(r2["run_id"])["state"] == "succeeded"
    await j.close()
    s.close()


@pytest.mark.asyncio
async def test_workflow_missing_account_is_explicit(tmp_path):
    s = Store(tmp_path / "db")
    j = Jobs(s)
    settings = Settings()
    w = Workflows(s, j, Tools(s, lambda: settings), Model(), lambda: settings)
    d = w.save(WorkflowDefinition(title="Fixture", template="email_assistant"))
    r = w.start(d["id"])
    with pytest.raises(ValueError, match="Not connected"):
        await w.execute(r["run_id"])
    assert w.run(r["run_id"])["state"] == "failed"
    await j.close()
    s.close()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "template",
    ["morning_briefing", "research_monitor", "project_coordination", "email_assistant"],
)
async def test_template_contracts_with_substituted_reads(
    tmp_path, monkeypatch, template
):
    from jarvis import config

    monkeypatch.setattr(config, "DATA", tmp_path)
    s = Store(tmp_path / "db")
    j = Jobs(s)
    settings = Settings()
    tools = Tools(s, lambda: settings)

    async def execute(*args):
        return {"ok": True, "data": {"items": [], "cache_age_seconds": 200}}

    monkeypatch.setattr(tools, "execute", execute)

    class ReadConnector:
        async def read(self, profile):
            return [{"title": "Fixture task or mail"}]

    w = Workflows(s, j, tools, Model(), lambda: settings, ReadConnector())
    d = w.save(
        WorkflowDefinition(title="Fixture", template=template, inputs={"format": "md"})
    )
    run = w.start(d["id"])
    await w.execute(run["run_id"])
    assert w.run(run["run_id"])["state"] == "succeeded"
    await j.close()
    s.close()


def test_strict_workflow_validation():
    for definition in [
        {"title": "Fixture", "template": "purchase"},
        {"title": "Fixture", "template": "voice_document", "enabled": "yes"},
        {
            "title": "Fixture",
            "template": "voice_document",
            "inputs": {"path": "/etc/passwd"},
        },
        {
            "title": "Fixture",
            "template": "voice_document",
            "schedule": {"minutes": 1, "kind": "interval"},
        },
    ]:
        with pytest.raises(ValueError):
            WorkflowDefinition.model_validate(definition)
