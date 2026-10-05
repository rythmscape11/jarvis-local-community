"""Bounded local document export. Never executes markup or fetches remote assets."""

import hashlib
import os
import re
from pathlib import Path
from xml.sax.saxutils import escape
from . import config
from .store import uid, now


def create_document(store, title, content, format):
    if format not in {"md", "pdf", "docx", "csv", "xlsx", "pptx"}:
        raise ValueError("Choose docx, pdf, md, csv, xlsx or pptx")
    if not title.strip() or not content.strip() or len(content) > 20000:
        raise ValueError("A document needs a bounded title and content")
    folder = config.DATA / "generated-documents"
    folder.mkdir(mode=0o700, parents=True, exist_ok=True)
    if folder.is_symlink() or folder.resolve().parent != config.DATA.resolve():
        raise ValueError("Document directory must remain inside private Jarvis data")
    id = uid()
    slug = re.sub(r"[^a-zA-Z0-9_-]+", "-", title).strip("-")[:60] or "document"
    target = folder / f"{slug}-{id}.{format}"
    temporary = target.with_suffix(".partial")
    lines = content.splitlines()
    try:
        if format == "md":
            temporary.write_text(
                "# " + title + "\n\n" + content + "\n\n---\nDraft for owner review.\n",
                encoding="utf-8",
            )
        elif format == "docx":
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor

            document = Document()
            section = document.sections[0]
            section.top_margin = section.bottom_margin = Inches(0.8)
            style = document.styles["Normal"]
            style.font.name = "Calibri"
            style.font.size = Pt(11)
            style.paragraph_format.space_after = Pt(8)
            document.add_heading(title, 0)
            for line in lines:
                value = line.strip()
                if not value:
                    continue
                heading = re.match(r"^(#{1,3})\s+(.+)", value)
                bullet = re.match(r"^[-*•]\s+(.+)", value)
                numbered = re.match(r"^\d+[.)]\s+(.+)", value)
                if heading:
                    document.add_heading(heading[2], len(heading[1]))
                elif bullet:
                    document.add_paragraph(bullet[1], style="List Bullet")
                elif numbered:
                    document.add_paragraph(numbered[1], style="List Number")
                else:
                    document.add_paragraph(value)
            footer = section.footer.paragraphs[0]
            run = footer.add_run("Draft for owner review · Jarvis Local")
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string("64706B")
            document.save(temporary)
            verified = Document(temporary)
            if not verified.paragraphs or verified.paragraphs[0].text != title:
                raise ValueError("Word document verification failed")
        elif format in {"csv", "xlsx"}:
            import csv
            import io

            rows = list(csv.reader(io.StringIO(content)))[:1000]
            rows = [
                [
                    ("'" + cell if cell.startswith(("=", "+", "-", "@")) else cell)[
                        :2000
                    ]
                    for cell in row[:50]
                ]
                for row in rows
            ]
            if format == "csv":
                with temporary.open("w", newline="", encoding="utf-8") as file:
                    csv.writer(file).writerows(rows)
            else:
                from openpyxl import Workbook, load_workbook
                from openpyxl.styles import Font, PatternFill

                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Draft"
                for row in rows:
                    sheet.append(row)
                sheet.freeze_panes = "A2"
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="203E34")
                sheet.auto_filter.ref = sheet.dimensions
                workbook.save(temporary)
                with temporary.open("rb") as file:
                    verified = load_workbook(file)
                    if verified.active.max_row != len(rows):
                        raise ValueError("Spreadsheet verification failed")
        elif format == "pptx":
            from pptx import Presentation
            from pptx.util import Inches, Pt

            presentation = Presentation()
            presentation.slide_width = Inches(13.333)
            presentation.slide_height = Inches(7.5)
            slide = presentation.slides.add_slide(presentation.slide_layouts[0])
            slide.shapes.title.text = title
            slide.placeholders[1].text = "Draft for owner review · Jarvis Local"
            chunks = []
            for line in lines:
                if line.strip():
                    chunks.append(line.strip()[:240])
            for offset in range(0, min(len(chunks), 174), 6):
                slide = presentation.slides.add_slide(presentation.slide_layouts[1])
                slide.shapes.title.text = title[:70] + f" · {offset // 6 + 1}"
                frame = slide.placeholders[1].text_frame
                for i, line in enumerate(chunks[offset : offset + 6]):
                    paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                    paragraph.text = line.lstrip("#-* ")
                    paragraph.font.size = Pt(22)
            if len(chunks) > 174:
                raise ValueError("Presentation exceeds 30 slides; shorten input")
            presentation.save(temporary)
            if len(Presentation(temporary).slides) < 2:
                raise ValueError("Presentation verification failed")
        else:
            if re.search(r"[\u0980-\u09ff]", title + content):
                raise ValueError(
                    "Bengali PDF font support is not configured; use DOCX or Markdown"
                )
            import reportlab
            from reportlab.lib import colors
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.enums import TA_LEFT
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.pagesizes import A4

            font = "JarvisVera"
            if font not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(
                    TTFont(
                        font, str(Path(reportlab.__file__).parent / "fonts/Vera.ttf")
                    )
                )
            styles = getSampleStyleSheet()
            for name in ["Normal", "Title", "Heading1", "Heading2", "Heading3"]:
                styles[name].fontName = font
                styles[name].alignment = TA_LEFT
            styles["Normal"].leading = 16
            styles["Normal"].fontSize = 10.5
            styles["Title"].textColor = colors.HexColor("#203E34")
            story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 18)]
            for line in lines:
                value = line.strip()
                if not value:
                    story.append(Spacer(1, 7))
                    continue
                heading = re.match(r"^(#{1,3})\s+(.+)", value)
                bullet = re.match(r"^[-*•]\s+(.+)", value)
                if heading:
                    story.append(
                        Paragraph(
                            escape(heading[2]), styles["Heading" + str(len(heading[1]))]
                        )
                    )
                elif bullet:
                    story.append(Paragraph("• " + escape(bullet[1]), styles["Normal"]))
                else:
                    story.append(Paragraph(escape(value), styles["Normal"]))
                story.append(Spacer(1, 6))
            story += [
                Spacer(1, 16),
                Paragraph("Draft for owner review · Jarvis Local", styles["Normal"]),
            ]
            SimpleDocTemplate(
                str(temporary),
                pagesize=A4,
                rightMargin=50,
                leftMargin=50,
                topMargin=50,
                bottomMargin=50,
                title=title,
            ).build(story)
            from pypdf import PdfReader

            reader = PdfReader(temporary)
            if not reader.pages or not reader.pages[0].extract_text().strip():
                raise ValueError("PDF verification failed")
        os.chmod(temporary, 0o600)
        temporary.replace(target)
        body = target.read_bytes()
        digest = hashlib.sha256(body).hexdigest()
        store.run(
            "INSERT INTO generated_documents VALUES(?,?,?,?,?,?,?)",
            (id, title, format, str(target), digest, len(body), now()),
        )
        return {
            "id": id,
            "title": title,
            "format": format,
            "filename": target.name,
            "bytes": len(body),
            "sha256": digest,
            "status": "completed",
            "download_url": f"/api/generated-documents/{id}/download",
            "review_required": True,
        }
    except BaseException:
        temporary.unlink(missing_ok=True)
        target.unlink(missing_ok=True)
        raise


def document_path(store, id):
    rows = store.all("SELECT * FROM generated_documents WHERE id=?", (id,))
    if not rows:
        raise ValueError("Document not found")
    target = Path(rows[0]["path"])
    folder = (config.DATA / "generated-documents").resolve()
    if target.is_symlink() or not target.is_file() or target.resolve().parent != folder:
        raise ValueError("Document is missing or outside the allowed directory")
    if hashlib.sha256(target.read_bytes()).hexdigest() != rows[0]["sha256"]:
        raise ValueError("Document changed since creation; regenerate or inspect it")
    return target
